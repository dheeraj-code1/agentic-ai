import pymupdf 
from docx import Document

import os



def read_pdf_file(file_path):
    if not os.path.exists(file_path):
        raise ValueError("file not found")
    with pymupdf.open(file_path) as doc:
        text = ""
        for page in doc:
            text += page.get_text()
    return text



def read_docx(file_path):
    if not os.path.exists(file_path):
        raise ValueError("file not found")
    doc = Document(file_path)

    content = []

    # Main document paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            content.append(text)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                row_data.append(cell.text.strip())

            content.append(" | ".join(row_data))

    # Headers and footers
    for section in doc.sections:

        # Header
        for paragraph in section.header.paragraphs:
            text = paragraph.text.strip()
            if text:
                content.append(text)

        # Footer
        for paragraph in section.footer.paragraphs:
            text = paragraph.text.strip()
            if text:
                content.append(text)

    return "\n".join(content)


def read_resume(path):
    file_type = path.split(".")[-1]

    if file_type=="pdf":
        return read_pdf_file(path)
    elif file_type=="docx":
        return read_docx(path)
    else:
        return None